from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Iterable

from .models import Document


SUPPORTED_EXTENSIONS = {".txt", ".md", ".json", ".docx"}


def discover_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    if not input_path.exists():
        raise FileNotFoundError(f"Input path does not exist: {input_path}")
    return sorted(
        path
        for path in input_path.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def load_documents(paths: Iterable[Path]) -> list[Document]:
    documents: list[Document] = []
    for path in paths:
        text = read_text(path)
        title = infer_title(path, text)
        body = strip_title_from_content(text, title)
        document_id = hashlib.sha1(str(path).encode("utf-8")).hexdigest()[:12]
        documents.append(
            Document(
                document_id=document_id,
                title=title,
                source_path=path,
                source_type=path.suffix.lower().lstrip("."),
                content=normalize_text(body),
                provider="local_file",
            )
        )
    return documents


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            return str(payload.get("content") or payload.get("text") or "")
        return json.dumps(payload, ensure_ascii=True, indent=2)
    if suffix == ".docx":
        try:
            from docx import Document as WordDocument
        except ImportError as exc:
            raise RuntimeError(
                "DOCX support requires python-docx. Install it or use TXT/MD input."
            ) from exc
        document = WordDocument(path)
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    raise ValueError(f"Unsupported document format: {path.suffix}")


def infer_title(path: Path, text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped[:120]
    return path.stem.replace("_", " ")


def normalize_text(text: str) -> str:
    normalized_lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in normalized_lines if line).strip()


def strip_title_from_content(text: str, title: str) -> str:
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if line.strip() == title.strip():
            remaining = lines[index + 1 :]
            return "\n".join(remaining) if remaining else text
    return text
